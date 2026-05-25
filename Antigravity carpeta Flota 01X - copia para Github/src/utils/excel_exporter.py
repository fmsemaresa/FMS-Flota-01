import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import joinedload
from src.database.connection import get_session
from src.database.models import Vehicle, Driver, UsageRecord

def export_db_to_excel(output_path):
    """Exporta los datos de SQLite a un archivo Excel bien formateado."""
    session = get_session()
    
    try:
        wb = openpyxl.Workbook()
        # Eliminar hoja por defecto
        default_sheet = wb.active
        wb.remove(default_sheet)
        
        # Estilos premium para cabeceras
        font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        fill_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid") # Azul oscuro
        alignment_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        alignment_left = Alignment(horizontal="left", vertical="center")
        
        border_thin = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )
        
        font_data = Font(name="Segoe UI", size=10)
        
        # ----------------------------------------------------
        # HOJA 1: HISTORIAL DE USOS
        # ----------------------------------------------------
        ws_usage = wb.create_sheet(title="Historial de Usos")
        headers_usage = [
            "ID Registro", "Patente Vehículo", "Modelo Vehículo", "Tipo Vehículo",
            "RUT Conductor", "Nombre Conductor", "Rol Conductor",
            "Fecha/Hora Inicio", "Fecha/Hora Fin", "Duración (m)", 
            "GPS Inicio", "GPS Fin", "Estado"
        ]
        
        ws_usage.append(headers_usage)
        
        records = (
            session.query(UsageRecord)
            .options(joinedload(UsageRecord.vehicle), joinedload(UsageRecord.driver))
            .order_by(UsageRecord.start_time.desc())
            .all()
        )
        for r in records:
            # Calcular duración en minutos si ya terminó
            duration = ""
            if r.end_time and r.start_time:
                diff = r.end_time - r.start_time
                duration = round(diff.total_seconds() / 60)
            
            start_str = r.start_time.strftime("%Y-%m-%d %H:%M:%S") if r.start_time else ""
            end_str = r.end_time.strftime("%Y-%m-%d %H:%M:%S") if r.end_time else "En uso activo"
            
            ws_usage.append([
                r.id,
                r.vehicle.plate,
                r.vehicle.model,
                r.vehicle.vehicle_type,
                r.driver.rut,
                r.driver.name,
                r.driver.role,
                start_str,
                end_str,
                duration,
                r.start_gps or "",
                r.end_gps or "",
                "Activo" if r.status == "active" else "Terminado"
            ])
            
        # ----------------------------------------------------
        # HOJA 2: VEHÍCULOS
        # ----------------------------------------------------
        ws_vehicles = wb.create_sheet(title="Vehículos")
        headers_vehicles = ["ID", "Patente", "Modelo", "Tipo Vehículo", "Estado Operativo"]
        ws_vehicles.append(headers_vehicles)
        
        vehicles = session.query(Vehicle).all()
        for v in vehicles:
            ws_vehicles.append([
                v.id,
                v.plate,
                v.model,
                v.vehicle_type,
                "Disponible" if v.is_active else "Inactivo"
            ])
            
        # ----------------------------------------------------
        # HOJA 3: CONDUCTORES
        # ----------------------------------------------------
        ws_drivers = wb.create_sheet(title="Conductores")
        headers_drivers = ["ID", "RUT", "Nombre Completo", "Rol", "Email", "Teléfono"]
        ws_drivers.append(headers_drivers)
        
        drivers = session.query(Driver).all()
        for d in drivers:
            ws_drivers.append([
                d.id,
                d.rut,
                d.name,
                d.role,
                d.email or "",
                d.phone or ""
            ])
            
        # Aplicar formato a todas las hojas creadas
        for ws in wb.worksheets:
            # Mostrar cuadrícula
            ws.views.sheetView[0].showGridLines = True
            
            # Formatear fila de cabecera
            for cell in ws[1]:
                cell.font = font_header
                cell.fill = fill_header
                cell.alignment = alignment_center
                
            # Establecer alto de fila de cabecera
            ws.row_dimensions[1].height = 25
            
            # Formatear filas de datos
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                for cell in row:
                    cell.font = font_data
                    cell.border = border_thin
                    # Alineaciones según tipo de dato
                    if isinstance(cell.value, int) or (cell.value and cell.value.isdigit()):
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    else:
                        cell.alignment = alignment_left
                        
            # Ajustar anchos de columnas automáticamente
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
                
        # Guardar archivo
        wb.save(output_path)
        return True
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        print(f"Error al exportar a Excel: {e}")
        return False
    finally:
        session.close()
